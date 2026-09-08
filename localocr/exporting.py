from __future__ import annotations

import html
import json
from pathlib import Path

from .core import table_shape


def table_html(table):
    rows, cols = table_shape(table)
    anchors = {(c["row"], c["col"]): c for c in table["cells"]}
    lines = ["<table>"]
    for row in range(rows):
        lines.append("<tr>")
        for col in range(cols):
            cell = anchors.get((row, col))
            if cell:
                lines.append(f'<td rowspan="{cell.get("rowspan", 1)}" colspan="{cell.get("colspan", 1)}">'
                             + html.escape(cell["text"]).replace("\n", "<br>") + "</td>")
        lines.append("</tr>")
    return "\n".join(lines + ["</table>"])


def render_text(doc, markdown=False):
    if doc.get('mode') == 'code':
        text = '\n'.join(b['text'] for b in doc['blocks'])
        if not markdown:
            return text
        import re
        length = max([len(m) for m in re.findall(r'`+', text)] + [2]) + 1
        fence = '`' * max(3, length)
        return fence + doc.get('code_options', {}).get('language', 'text') + '\n' + text + '\n' + fence
    tables = {t["id"]: t for t in doc["tables"]}
    parts, used = [], set()
    for b in doc["blocks"]:
        if b.get("table_id") in tables:
            t = tables[b["table_id"]]
            used.add(t["id"])
            parts.append(table_html(t) if markdown else plain_table(t))
        else:
            text = b["text"]
            if markdown and b["kind"] in {"title", "doc_title", "paragraph_title"}:
                text = "## " + text
            parts.append(text)
    for t in doc["tables"]:
        if t["id"] not in used:
            parts.append(table_html(t) if markdown else plain_table(t))
    return ("\n\n" if markdown else "\n").join(parts)


def plain_table(table):
    rows, cols = table_shape(table)
    grid = [[""] * cols for _ in range(rows)]
    for cell in table["cells"]:
        grid[cell["row"]][cell["col"]] = cell["text"].replace("\n", " ")
    return "\n".join("\t".join(row) for row in grid)


def export_project(store, root, output, reviewed_only=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    combined_txt, combined_md, count = [], [], 0
    for page in store.pages(root):
        if reviewed_only and not page["reviewed"]:
            continue
        doc = store.document(page["id"])
        if doc is None:
            continue
        relative = Path(page["relative"])
        if relative.is_absolute() or ".." in relative.parts:
            raise ValueError("无效的导出相对路径")
        # Preserve extension too: scan.jpg and scan.png must never overwrite each other.
        base = output / relative.parent / f"{relative.name}.page-{page['page']+1:04d}"
        base.parent.mkdir(parents=True, exist_ok=True)
        txt, md = render_text(doc), render_text(doc, True)
        if doc.get('mode') == 'code':
            from .codeocr import LANGUAGES
            extension = LANGUAGES.get(doc.get('code_options', {}).get('language'), '.txt')
            Path(str(base) + '.source' + extension).write_text(txt, encoding='utf-8', newline='')
        Path(str(base) + ".txt").write_text(txt, encoding="utf-8-sig")
        Path(str(base) + ".md").write_text(md, encoding="utf-8")
        versions = [{**v, "original": store.version_document(v["id"], original=True)} for v in store.versions(page["id"])]
        Path(str(base) + ".json").write_text(json.dumps({"page": page, "current": doc, "versions": versions}, ensure_ascii=False, indent=2), encoding="utf-8")
        if doc["tables"]:
            workbook = Workbook()
            workbook.remove(workbook.active)
            for index, table in enumerate(doc["tables"], 1):
                sheet = workbook.create_sheet(f"表格{index}")
                for cell in table["cells"]:
                    row, col = cell["row"]+1, cell["col"]+1
                    target = sheet.cell(row, col, cell["text"])
                    target.data_type = "s"  # OCR text beginning '=' is text, never a formula.
                    target.alignment = Alignment(vertical="top", wrap_text=True)
                    if row == 1:
                        target.font = Font(bold=True, color="FFFFFF")
                        target.fill = PatternFill("solid", fgColor="22695C")
                    if cell.get("rowspan", 1) > 1 or cell.get("colspan", 1) > 1:
                        sheet.merge_cells(start_row=row, start_column=col,
                                          end_row=row+cell.get("rowspan", 1)-1,
                                          end_column=col+cell.get("colspan", 1)-1)
                for column in sheet.columns:
                    from openpyxl.utils import get_column_letter
                    sheet.column_dimensions[get_column_letter(column[0].column)].width = 24
                sheet.freeze_panes = "A2"
            workbook.save(str(base) + ".xlsx")
        title = f"{page['relative']} · 第 {page['page']+1} 页"
        combined_txt.append(title + "\n" + txt)
        combined_md.append("# " + title + "\n\n" + md)
        count += 1
    (output / "全部结果.txt").write_text("\n\n".join(combined_txt), encoding="utf-8-sig")
    (output / "全部结果.md").write_text("\n\n---\n\n".join(combined_md), encoding="utf-8")
    return count
